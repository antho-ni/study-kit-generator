# backend/app/services/extraction.py
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
import magic
import io

class ExtractionError(Exception):
    pass

# Maps each accepted extension to the real MIME types libmagic should detect
EXPECTED_MIME_TYPES = {
    ".pdf": ["application/pdf"],
    ".docx": ["application/zip", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"],
    ".pptx": ["application/zip", "application/vnd.openxmlformats-officedocument.presentationml.presentation"],
    ".txt": ["text/plain"],
}

def _verify_file_type(filename: str, file_bytes: bytes) -> None:
    """Confirm the file's actual content matches its claimed extension,
    rather than trusting the filename alone."""
    ext = Path(filename).suffix.lower()
    detected_mime = magic.from_buffer(file_bytes, mime=True)

    expected = EXPECTED_MIME_TYPES.get(ext, [])
    if detected_mime not in expected:
        raise ExtractionError(
            f"This file's content doesn't match a valid {ext} file. "
            f"It may be renamed, corrupted, or a different file type."
        )

def extract_text(filename: str, file_bytes: bytes) -> str:
    ext = Path(filename).suffix.lower()

    if ext not in EXPECTED_MIME_TYPES:
        raise ExtractionError(f"Unsupported file type: {ext}")

    _verify_file_type(filename, file_bytes)

    if ext == ".pdf":
        return _extract_pdf(file_bytes)
    elif ext == ".docx":
        return _extract_docx(file_bytes)
    elif ext == ".pptx":
        return _extract_pptx(file_bytes)
    elif ext == ".txt":
        return file_bytes.decode("utf-8", errors="ignore")

def _extract_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    text_parts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            text_parts.append(text)
    full_text = "\n\n".join(text_parts)
    if not full_text.strip():
        raise ExtractionError(
            "No extractable text found — this PDF may be scanned/image-based and needs OCR."
        )
    return full_text

def _extract_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text for cell in row.cells)
            if row_text.strip():
                paragraphs.append(row_text)
    return "\n".join(paragraphs)

def _extract_pptx(file_bytes: bytes) -> str:
    prs = Presentation(io.BytesIO(file_bytes))
    slides_text = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_text = [f"--- Slide {i} ---"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs)
                    if text.strip():
                        slide_text.append(text)
        slides_text.append("\n".join(slide_text))
    return "\n\n".join(slides_text)